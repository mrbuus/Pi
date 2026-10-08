from pathlib import Path
import openpyxl,json,datetime,hashlib,collections,argparse
parser=argparse.ArgumentParser(description='Read-only local schedule source audit; never changes workbook')
parser.add_argument('--source',default='/Users/mr.buus/Downloads/Хуваарь.xlsx')
args=parser.parse_args()
p=Path(args.source);out=Path(__file__).resolve().parent/'private';out.mkdir(exist_ok=True)
w=openpyxl.load_workbook(p,data_only=False);v=openpyxl.load_workbook(p,data_only=True)
def val(x):return x.isoformat() if isinstance(x,(datetime.datetime,datetime.date,datetime.time)) else x
def color(c):return str(c.rgb) if c and c.type=='rgb' else (str(c.type)+':'+str(c.index) if c else None)
a={'source':str(p),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'properties':{'created':val(w.properties.created),'modified':val(w.properties.modified)},'sheets':[]}
for s in w:
 cells=[]
 for row in s:
  for c in row:
   if c.value is not None:
    cells.append({'cell':c.coordinate,'r':c.row,'c':c.column,'value':val(c.value),'cached':val(v[s.title][c.coordinate].value),'type':c.data_type,'format':c.number_format,'fill':color(c.fill.fgColor),'fontColor':color(c.font.color),'bold':c.font.bold,'comment':c.comment.text if c.comment else None})
 a['sheets'].append({'name':s.title,'state':s.sheet_state,'rows':s.max_row,'cols':s.max_column,'merged':[str(x) for x in s.merged_cells.ranges],'cells':cells})
 print(s.title, s.max_row,s.max_column,'cells',len(cells),'merges',len(s.merged_cells.ranges),'formula',sum(c['type']=='f' for c in cells))
 print(' / '.join(str(c['value'])[:95] for c in cells[:20]))
(out/'schedule-audit.json').write_text(json.dumps(a,ensure_ascii=False,indent=1))
print('properties',a['properties'])
