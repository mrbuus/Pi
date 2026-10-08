"""Read supplied workbooks; preserve a private ID snapshot; never access a DB."""
from pathlib import Path
import collections, datetime, hashlib, json, re, shutil, unicodedata
import openpyxl

HERE = Path(__file__).resolve().parent
BASE = Path('/Users/mr.buus/Downloads')
PRIVATE = HERE / 'private'
PRIVATE.mkdir(exist_ok=True)

def norm(value):
    return re.sub(r'\s+', ' ', unicodedata.normalize('NFKC', str(value or ''))).strip().casefold()

def ident(value):
    if isinstance(value, (int, float)) and value == int(value):
        return str(int(value))
    return str(value or '').strip()

def phone(value):
    text = re.sub(r'\.0$', '', ident(value))
    return set(re.findall(r'(?<!\d)(?:\+?976[ -]?)?([5-9]\d{7})(?!\d)', text))

registry_file = BASE / 'Burtgel.xlsx'
workbook = openpyxl.load_workbook(registry_file, data_only=False)
sheet = workbook['Code']
registry = []
for row in sheet.iter_rows(min_row=2, values_only=True):
    if not any(v is not None for v in row):
        continue
    registry.append(dict(zip(['newCode','lastName','firstName','grade','studentPhone','fatherPhone','motherPhone','school','section'], [ident(v) for v in row[:9]])))
codes = collections.Counter(r['newCode'] for r in registry)
assert all(v == 1 for v in codes.values()), 'Duplicate source ID; do not map automatically.'
assert all(re.fullmatch(r'\d{7}', k) for k in codes), 'Unexpected ID format.'
shutil.copyfile(registry_file, PRIVATE / 'Burtgel-source.xlsx')

old_file = HERE.parents[1] / 'api/prisma/reports/student-import.json'
old = json.loads(old_file.read_text()).get('plannedStudents', []) if old_file.exists() else []
mapping = []
for r in registry:
    named = [p for p in old if (norm(p.get('lastName')),norm(p.get('firstName'))) == (norm(r['lastName']),norm(r['firstName']))]
    rp = set().union(*(phone(r[k]) for k in ['studentPhone','fatherPhone','motherPhone']))
    matched = [p for p in named if rp & set().union(*(phone(p.get(k)) for k in ['studentPhone','fatherPhone','motherPhone']))]
    state = 'NAME_AND_PHONE_CANDIDATE' if len(matched)==1 else 'AMBIGUOUS' if len(matched)>1 else 'NAME_ONLY_REVIEW' if named else 'NO_OLD_REPORT_MATCH'
    mapping.append({'newCode':r['newCode'],'status':state,'registry':r,'oldReportRows':[{'sheet':p.get('sheet'),'row':p.get('row'),'key':p.get('key')} for p in matched or named], 'databaseUserId':None,'oldStudentCode':None})
(PRIVATE / 'id-mapping-candidates.json').write_text(json.dumps(mapping,ensure_ascii=False,indent=2))
payment_file = next(p for p in BASE.glob('2027*.xlsx') if 'төлбөр' in p.name)
payments = openpyxl.load_workbook(payment_file, data_only=False)
cached = openpyxl.load_workbook(payment_file, data_only=True)
stats=[]; names=collections.Counter(); errors=[]
for s in payments:
    count=formulas=populated=0; header=[]; amount_counts=collections.Counter()
    for row in s:
        occupied=False
        for c in row:
            if c.value is None: continue
            occupied=True;count+=1
            if c.row <= 2: header.append({'cell':c.coordinate,'value':str(c.value)})
            if c.data_type=='f': formulas+=1
            if c.data_type=='e' or cached[s.title][c.coordinate].data_type=='e': errors.append({'sheet':s.title,'cell':c.coordinate})
            if isinstance(c.value,(int,float)) and c.value in [300000,350000,1400000,2500000,2850000]:amount_counts[int(c.value)]+=1
        populated+=occupied
        if s.title!='Шийт29' and row[1].row>2 and row[1].value and row[2].value:
            names[(norm(row[1].value),norm(row[2].value))]+=1
    stats.append({'sheet':s.title,'extent':[s.max_row,s.max_column],'nonemptyRows':populated,'cells':count,'formulas':formulas,'amountOccurrences':dict(amount_counts),'header':header})
summary={'date':'2026-09-22','scope':'read-only source audit; no database query or mutation', 'registry':{'rows':len(registry),'uniqueIds':len(codes),'duplicateIds':sum(v>1 for v in codes.values()),'grades':dict(collections.Counter(norm(r['grade']) for r in registry)),'sameFullNameGroups':sum(v>1 for v in collections.Counter((norm(r['lastName']),norm(r['firstName'])) for r in registry).values()),'sha256':hashlib.sha256(registry_file.read_bytes()).hexdigest(),'oldReportRows':len(old),'mappingStates':dict(collections.Counter(m['status'] for m in mapping))},'payments':{'sheets':stats,'formulaErrors':errors,'repeatedNameGroupsAcrossRows':sum(v>1 for v in names.values()),'warning':'Repeated names are candidates, not proven duplicate students or payments.'}}
(PRIVATE / 'source-audit.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2))
print(json.dumps({'registry':summary['registry'],'sheets':[{k:v for k,v in s.items() if k!='header'} for s in stats],'formulaErrors':errors,'repeatedNameGroups':summary['payments']['repeatedNameGroupsAcrossRows']},ensure_ascii=False,indent=2))
